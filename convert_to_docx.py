#!/usr/bin/env python3
"""
Workmate Documentation to DOCX Converter

This script converts MkDocs documentation to a single DOCX file.
It reads the mkdocs.yml configuration and processes all markdown files
according to the navigation structure.
"""

import os
import re
import sys
import requests
import io
from pathlib import Path
from typing import Dict, List, Optional, Any
from urllib.parse import urlparse, urljoin

import yaml
import markdown
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from bs4 import BeautifulSoup
from PIL import Image


class MkDocsToDocxConverter:
    """Converts MkDocs documentation to a single DOCX file."""

    def __init__(self, project_root: Path):
        self.project_root = project_root
        self.docs_dir = project_root / "docs"
        self.mkdocs_config = self._load_mkdocs_config()
        self.document = Document()
        self.md = markdown.Markdown(extensions=[
            'extra',
            'codehilite',
            'toc',
            'tables',
            'fenced_code'
        ])
        self._setup_document_styles()

    def _load_mkdocs_config(self) -> Dict[str, Any]:
        """Load and parse the mkdocs.yml configuration file."""
        config_file = self.project_root / "mkdocs.yml"
        if not config_file.exists():
            raise FileNotFoundError(f"mkdocs.yml not found at {config_file}")
        
        with open(config_file, 'r', encoding='utf-8') as f:
            return yaml.safe_load(f)

    def _setup_document_styles(self):
        """Set up custom styles for the document."""
        styles = self.document.styles

        # Title style
        if 'Document Title' not in [s.name for s in styles]:
            title_style = styles.add_style('Document Title', WD_STYLE_TYPE.PARAGRAPH)
            title_format = title_style.paragraph_format
            title_format.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            title_format.space_after = Pt(24)
            title_font = title_style.font
            title_font.name = 'Arial'
            title_font.size = Pt(24)
            title_font.bold = True

        # Section header style
        if 'Section Header' not in [s.name for s in styles]:
            section_style = styles.add_style('Section Header', WD_STYLE_TYPE.PARAGRAPH)
            section_format = section_style.paragraph_format
            section_format.space_before = Pt(18)
            section_format.space_after = Pt(12)
            section_font = section_style.font
            section_font.name = 'Arial'
            section_font.size = Pt(16)
            section_font.bold = True

        # Code style
        if 'Code Block' not in [s.name for s in styles]:
            code_style = styles.add_style('Code Block', WD_STYLE_TYPE.PARAGRAPH)
            code_format = code_style.paragraph_format
            code_format.space_before = Pt(6)
            code_format.space_after = Pt(6)
            code_font = code_style.font
            code_font.name = 'Courier New'
            code_font.size = Pt(10)

    def _add_title_page(self):
        """Add a title page to the document."""
        # Title
        title = self.document.add_paragraph(
            self.mkdocs_config.get('site_name', 'Documentation'),
            style='Document Title'
        )
        
        # Subtitle
        subtitle = self.document.add_paragraph(
            'Complete Documentation Export',
            style='Normal'
        )
        subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        # Add some space
        self.document.add_paragraph()
        
        # Generation info
        from datetime import datetime
        info = self.document.add_paragraph(
            f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            style='Normal'
        )
        info.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        # Page break
        self.document.add_page_break()

    def _add_table_of_contents(self):
        """Add a table of contents based on navigation structure."""
        toc_title = self.document.add_paragraph("Table of Contents", style='Section Header')
        
        def add_nav_items(nav_items: List, level: int = 0):
            for item in nav_items:
                if isinstance(item, dict):
                    for title, content in item.items():
                        indent = "  " * level
                        if isinstance(content, str):
                            # Single page
                            toc_item = self.document.add_paragraph(f"{indent}• {title}")
                        elif isinstance(content, list):
                            # Section with sub-items
                            toc_item = self.document.add_paragraph(f"{indent}• {title}")
                            add_nav_items(content, level + 1)
                elif isinstance(item, str):
                    # Simple item
                    indent = "  " * level
                    toc_item = self.document.add_paragraph(f"{indent}• {item}")

        nav = self.mkdocs_config.get('nav', [])
        add_nav_items(nav)
        
        # Page break
        self.document.add_page_break()

    def _convert_markdown_to_docx_elements(self, md_content: str, file_path: Path):
        """Convert markdown content to DOCX elements."""
        # Convert markdown to HTML
        html = self.md.convert(md_content)
        soup = BeautifulSoup(html, 'html.parser')
        
        # Process each top-level element
        for element in soup.children:
            if hasattr(element, 'name') and element.name:
                if element.name in ['h1', 'h2', 'h3', 'h4', 'h5', 'h6']:
                    self._add_heading(element)
                elif element.name == 'p':
                    self._add_paragraph(element)
                elif element.name == 'ul':
                    self._add_unordered_list(element)
                elif element.name == 'ol':
                    self._add_ordered_list(element)
                elif element.name == 'table':
                    self._add_table(element)
                elif element.name == 'pre':
                    self._add_code_block(element)
                elif element.name == 'blockquote':
                    self._add_blockquote(element)
                elif element.name == 'img':
                    self._add_image(element)

    def _add_heading(self, element):
        """Add a heading to the document."""
        level = int(element.name[1])  # h1 -> 1, h2 -> 2, etc.
        text = element.get_text().strip()
        
        if text:  # Only add non-empty headings
            if level == 1:
                heading = self.document.add_paragraph(text, style='Heading 1')
            elif level == 2:
                heading = self.document.add_paragraph(text, style='Heading 2')
            elif level == 3:
                heading = self.document.add_paragraph(text, style='Heading 3')
            else:
                heading = self.document.add_paragraph(text, style='Heading 4')

    def _add_paragraph(self, element):
        """Add a paragraph to the document."""
        # Check if paragraph contains images
        images = element.find_all('img')
        if images:
            # Handle paragraphs with images
            for img in images:
                self._add_image(img)
            # Also add text content if any
            text = element.get_text().strip()
            if text:
                paragraph = self.document.add_paragraph(text)
        else:
            # Handle regular paragraphs
            text = element.get_text().strip()
            if text:
                paragraph = self.document.add_paragraph()
                
                # Handle inline formatting
                for content in element.contents:
                    if hasattr(content, 'name') and content.name:
                        if content.name == 'strong' or content.name == 'b':
                            run = paragraph.add_run(content.get_text())
                            run.bold = True
                        elif content.name == 'em' or content.name == 'i':
                            run = paragraph.add_run(content.get_text())
                            run.italic = True
                        elif content.name == 'code':
                            run = paragraph.add_run(content.get_text())
                            run.font.name = 'Courier New'
                        elif content.name == 'a':
                            # Add link text with note about original URL
                            link_text = content.get_text()
                            url = content.get('href', '')
                            if url:
                                run = paragraph.add_run(f"{link_text} ({url})")
                            else:
                                run = paragraph.add_run(link_text)
                        elif content.name == 'img':
                            # Skip images here as they're handled separately
                            pass
                        else:
                            paragraph.add_run(content.get_text())
                    else:
                        paragraph.add_run(str(content))

    def _add_unordered_list(self, element):
        """Add an unordered list to the document."""
        for li in element.find_all('li', recursive=False):
            text = li.get_text().strip()
            if text:
                paragraph = self.document.add_paragraph(text, style='List Bullet')

    def _add_ordered_list(self, element):
        """Add an ordered list to the document."""
        for li in element.find_all('li', recursive=False):
            text = li.get_text().strip()
            if text:
                paragraph = self.document.add_paragraph(text, style='List Number')

    def _add_table(self, element):
        """Add a table to the document."""
        rows = element.find_all('tr')
        if not rows:
            return
        
        # Get dimensions
        max_cols = max(len(row.find_all(['td', 'th'])) for row in rows) if rows else 0
        if max_cols == 0:
            return
        
        # Create table
        table = self.document.add_table(rows=len(rows), cols=max_cols)
        table.style = 'Table Grid'
        
        for i, row in enumerate(rows):
            cells = row.find_all(['td', 'th'])
            for j, cell in enumerate(cells):
                if j < max_cols:
                    table.cell(i, j).text = cell.get_text().strip()

    def _add_code_block(self, element):
        """Add a code block to the document."""
        code_text = element.get_text()
        if code_text.strip():
            paragraph = self.document.add_paragraph(code_text, style='Code Block')

    def _add_blockquote(self, element):
        """Add a blockquote to the document."""
        text = element.get_text().strip()
        if text:
            paragraph = self.document.add_paragraph(f'"{text}"')
            paragraph.style = 'Quote'

    def _add_image(self, element):
        """Add an image to the document."""
        src = element.get('src', '')
        alt = element.get('alt', '')
        
        if not src:
            # No image source, add alt text instead
            if alt:
                paragraph = self.document.add_paragraph(f"[Image: {alt}]")
            return
        
        try:
            # Determine if it's a URL or local file
            parsed_url = urlparse(src)
            
            if parsed_url.scheme in ['http', 'https']:
                # Download image from URL
                print(f"    📷 Downloading image: {src}")
                response = requests.get(src, timeout=10)
                response.raise_for_status()
                image_data = io.BytesIO(response.content)
            else:
                # Local file path
                if src.startswith('/'):
                    # Absolute path from docs root
                    image_path = self.docs_dir / src[1:]
                else:
                    # Relative path
                    image_path = self.docs_dir / src
                
                print(f"    📷 Loading local image: {image_path}")
                
                if not image_path.exists():
                    print(f"    ⚠️  Image not found: {image_path}")
                    paragraph = self.document.add_paragraph(f"[Image not found: {src}]")
                    return
                
                image_data = str(image_path)
            
            # Add image to document
            paragraph = self.document.add_paragraph()
            run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
            
            # Try to add the image with a reasonable size
            try:
                if isinstance(image_data, str):
                    # Local file
                    paragraph.add_run().add_picture(image_data, width=Inches(5))
                else:
                    # Downloaded image
                    paragraph.add_run().add_picture(image_data, width=Inches(5))
                
                # Add caption if alt text exists
                if alt:
                    caption_para = self.document.add_paragraph(f"Figure: {alt}")
                    caption_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                    
            except Exception as img_error:
                print(f"    ❌ Error adding image: {img_error}")
                # Fallback: add alt text or filename
                if alt:
                    paragraph.add_run(f"[Image: {alt}]")
                else:
                    paragraph.add_run(f"[Image: {src}]")
        
        except Exception as e:
            print(f"    ❌ Error processing image {src}: {e}")
            # Add fallback text
            paragraph = self.document.add_paragraph()
            if alt:
                paragraph.add_run(f"[Image: {alt} - {src}]")
            else:
                paragraph.add_run(f"[Image: {src}]")

    def _process_markdown_file(self, file_path: Path, title: str):
        """Process a single markdown file and add it to the document."""
        print(f"  📄 Processing: {title}")
        
        if not file_path.exists():
            print(f"    ⚠️  Warning: File not found: {file_path}")
            return

        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()

            # Add section header
            self.document.add_paragraph(title, style='Section Header')
            
            # Convert and add content
            self._convert_markdown_to_docx_elements(content, file_path)
            
            # Add some space after section
            self.document.add_paragraph()
            
        except Exception as e:
            print(f"    ❌ Error processing {file_path}: {e}")
            # Add error note to document
            error_para = self.document.add_paragraph(f"Error loading content from {file_path}: {e}")

    def _extract_nav_items(self, nav_items: List) -> List[tuple]:
        """Extract navigation items as (title, file_path) tuples."""
        items = []
        
        for item in nav_items:
            if isinstance(item, dict):
                for title, content in item.items():
                    if isinstance(content, str):
                        # Single page
                        items.append((title, content))
                    elif isinstance(content, list):
                        # Section with sub-items
                        items.extend(self._extract_nav_items(content))
            elif isinstance(item, str):
                # Simple format: just filename
                title = item.replace('.md', '').replace('_', ' ').title()
                items.append((title, item))
        
        return items

    def convert(self, output_file: str = "workmate_documentation.docx"):
        """Convert the MkDocs site to a DOCX file."""
        print("🚀 Starting conversion to DOCX...")
        
        # Add title page
        print("📋 Adding title page...")
        self._add_title_page()
        
        # Add table of contents
        print("📑 Adding table of contents...")
        self._add_table_of_contents()
        
        # Process navigation items
        nav = self.mkdocs_config.get('nav', [])
        nav_items = self._extract_nav_items(nav)
        
        print(f"📚 Processing {len(nav_items)} pages...")
        
        for title, file_path in nav_items:
            # Resolve file path
            if file_path.startswith('/'):
                full_path = self.docs_dir / file_path[1:]
            else:
                full_path = self.docs_dir / file_path
            
            self._process_markdown_file(full_path, title)
        
        # Save document
        output_path = self.project_root / output_file
        print(f"💾 Saving document to: {output_path}")
        self.document.save(output_path)
        print(f"✅ Conversion complete! Output saved to: {output_path}")
        
        return output_path


def main():
    """Main entry point for the script."""
    print("🔧 Workmate Documentation to DOCX Converter")
    print("=" * 50)
    
    # Get project root
    project_root = Path(__file__).parent
    print(f"🔍 Project root: {project_root}")
    
    # Parse command line arguments
    if len(sys.argv) > 1:
        output_file = sys.argv[1]
    else:
        output_file = "workmate_documentation.docx"
    
    print(f"📝 Output file: {output_file}")
    
    try:
        # Create converter and run conversion
        print("🚀 Initializing converter...")
        converter = MkDocsToDocxConverter(project_root)
        print("✅ Converter initialized successfully")
        
        output_path = converter.convert(output_file)
        
        print(f"\n🎉 Success! Your documentation has been converted to: {output_path}")
        if output_path.exists():
            print(f"📊 File size: {output_path.stat().st_size / (1024*1024):.1f} MB")
        
    except Exception as e:
        print(f"❌ Error during conversion: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == "__main__":
    main()
